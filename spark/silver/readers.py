"""Read successful immutable Bronze objects through verified streaming downloads.

Local Spark is intentional: the HTTP runner and local[*] share this temporary
file. Remote executors require a shared filesystem/S3A reader, not this path.
"""

import csv
import hashlib
import json
import os
import re
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import urlparse
from .registry import SOURCES


def client():
    from minio import Minio

    return Minio(
        os.environ.get("MINIO_ENDPOINT", "minio:9000"),
        access_key=os.environ["AWS_ACCESS_KEY_ID"],
        secret_key=os.environ["AWS_SECRET_ACCESS_KEY"],
        secure=os.environ.get("MINIO_SECURE", "false").lower() == "true",
    )


def resolve(dataset, connection=None):
    if dataset not in SOURCES:
        raise ValueError("Unknown dataset")
    c = connection or client()
    bucket = os.environ.get("BRONZE_BUCKET", "bronze")
    response = c.get_object(bucket, f"_metadata/latest_success/{dataset}.json")
    try:
        m = json.load(response)
    finally:
        response.close()
        response.release_conn()
    required = [
        "dataset_id",
        "ingestion_id",
        "checksum_sha256",
        "bronze_path",
        "source_system",
        "source_file_name",
    ]
    if (
        any(not m.get(k) for k in required)
        or m["dataset_id"] != dataset
        or m.get("status") != "SUCCESS"
    ):
        raise ValueError("Invalid Bronze success metadata")
    if not re.fullmatch("[a-f0-9]{64}", m["checksum_sha256"]):
        raise ValueError("Invalid Bronze checksum")
    u = urlparse(m["bronze_path"])
    if (
        u.scheme != "s3"
        or u.netloc != bucket
        or not u.path.lstrip("/")
        or u.query
        or u.fragment
    ):
        raise ValueError("Invalid Bronze object URI")
    if m.get("file_format") != SOURCES[dataset]["format"]:
        raise ValueError("Bronze format mismatch")
    if m["source_system"] != SOURCES[dataset]["source_system"]:
        raise ValueError("Source identity mismatch")
    stat = c.stat_object(bucket, u.path.lstrip("/"))
    if m.get("file_size_bytes") != stat.size:
        raise ValueError("Bronze object size mismatch")
    return m


@contextmanager
def download(metadata, connection=None):
    c = connection or client()
    u = urlparse(metadata["bronze_path"])
    with TemporaryDirectory(prefix="silver-") as directory:
        path = Path(directory) / ("input." + SOURCES[metadata["dataset_id"]]["format"])
        response = c.get_object(u.netloc, u.path.lstrip("/"))
        digest = hashlib.sha256()
        try:
            with path.open("wb") as out:
                for chunk in response.stream(1024 * 1024):
                    digest.update(chunk)
                    out.write(chunk)
        finally:
            response.close()
            response.release_conn()
        if digest.hexdigest() != metadata["checksum_sha256"]:
            raise ValueError("Bronze checksum mismatch")
        yield path


def iter_workbook(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.worksheets) != 1:
            raise ValueError("Unexpected workbook sheets")
        rows = wb.worksheets[0].iter_rows(values_only=True)
        for r in rows:
            yield ["" if v is None else str(v) for v in r]
    finally:
        wb.close()


def prepare(path, spec):
    """Validate exact input header before Spark; convert bounded XLSX to CSV."""
    path = Path(path)
    if spec["format"] == "xlsx":
        if path.stat().st_size > 20 * 1024 * 1024:
            raise ValueError("XLSX exceeds 20 MiB policy")
        destination = path.with_suffix(".csv")
        rows = iter_workbook(path)
        with destination.open("w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            for r in rows:
                writer.writerow(r)
        path = destination
    with path.open(encoding="utf-8-sig", newline="") as f:
        header = next(csv.reader(f))
    if header != spec["headers"]:
        raise ValueError("Input schema drift: reviewed header required")
    return path


def read(spark, path, spec):
    from pyspark.sql import types as T, functions as F

    path = prepare(path, spec)
    schema = T.StructType([T.StructField(h, T.StringType()) for h in spec["headers"]])
    frame = (
        spark.read.schema(schema)
        .options(
            header=True,
            enforceSchema=False,
            mode="FAILFAST",
            escape='"',
            multiLine=spec["target"] != "weather_daily",
            encoding="UTF-8",
        )
        .csv(str(path))
    )
    counts = {"input_rows": frame.count(), "excluded_rows": 0}
    if "source_sheet" in spec["headers"]:
        from .common import RULES

        monthly = frame.filter(F.col("source_sheet") == "Monthly Prices")
        # Only two small metadata rows are brought to the driver.
        labels = monthly.filter(
            F.col("column_12") == RULES["pink"]["column_12"][0]
        ).take(2)
        units = monthly.filter(F.col("column_12") == "($/kg)").take(2)
        if len(labels) != 1 or len(units) != 1:
            raise ValueError("Missing/ambiguous Pink Sheet header or unit row")
        for field, (name, _, _) in RULES["pink"].items():
            if labels[0][field] != name:
                raise ValueError("Pink Sheet column semantics changed")
        spec["pink_units"] = {field: units[0][field] for field in RULES["pink"]}
        frame = monthly.filter(F.col("column_0").rlike(r"^\d{4}M\d{2}$"))
        counts["excluded_rows"] = counts["input_rows"] - frame.count()
    if "motCode" in spec["headers"]:
        from .common import RULES

        rule = RULES["comtrade"]
        total = (
            (F.col("motCode") == rule["total_transport"])
            & (F.col("partner2Code") == rule["partner2"])
            & (F.col("customsCode") == rule["customs"])
        )
        keys = ["reporterCode", "partnerCode", "cmdCode", "period", "flowCode"]
        totals = frame.filter(total)
        breakdown = frame.filter(~total)
        # Only exclude detailed rows when a corresponding source total exists.
        covered = breakdown.join(totals.select(*keys).distinct(), keys, "left_semi")
        unmatched = breakdown.join(totals.select(*keys).distinct(), keys, "left_anti")
        counts["excluded_transport_rows"] = covered.count()
        counts["excluded_rows"] += counts["excluded_transport_rows"]
        frame = totals.unionByName(unmatched)
    return frame, counts
